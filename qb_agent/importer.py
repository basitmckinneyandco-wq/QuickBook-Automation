"""Turn an uploaded .csv / .xlsx file into clean WorkRecord objects.

Spreadsheets are messy: headers are spelled differently every time, dates come
in five formats, and money arrives as "$1,200.00" or "(200)". Everything in
this file exists to absorb that mess so the rest of the app only ever sees
tidy, typed values.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Mapping, Sequence

from .models import WorkRecord, deterministic_source_id

MAX_ROWS = 2000

# Every spelling we accept for each field. Compare against normalised headers:
# lowercased, punctuation turned into spaces, runs of spaces collapsed.
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "customer_name": (
        "customer", "customer name", "customer job", "client", "client name",
        "name", "full name", "corporation name", "bill to", "customer full name",
    ),
    "job_name": (
        "job", "job name", "policy", "policy name", "policy type", "line of business",
        "lob", "coverage", "coverage type",
    ),
    "company_name": ("company", "company name", "business name", "organization", "organisation"),
    "email": ("email", "e mail", "email address", "customer email"),
    "phone": ("phone", "phone number", "telephone", "contact number", "mobile"),
    "invoice_number": (
        "invoice number", "invoice no", "invoice", "invoice ref", "invoice num",
        "ref number", "reference", "reference number", "policy no", "policy number",
    ),
    "invoice_date": (
        "invoice date", "date", "txn date", "transaction date", "eff",
        "effective date", "service date",
    ),
    "item": ("item", "service", "product", "item name", "service item", "product service", "carrier"),
    "description": ("description", "desc", "memo", "line description", "details", "notes"),
    "quantity": ("quantity", "qty", "units", "hours"),
    "rate": ("rate", "unit price", "price", "unit rate", "hourly rate", "unit cost"),
    "amount": ("amount", "total", "line total", "invoice amount", "commission", "commision"),
}

DATE_FORMATS = (
    "%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y",
    "%m-%d-%Y", "%d %b %Y", "%b %d %Y", "%d %B %Y", "%B %d %Y", "%Y/%m/%d",
)


@dataclass(frozen=True, slots=True)
class ImportResult:
    records: list[WorkRecord]
    warnings: list[str]
    mapped_columns: dict[str, str]
    skipped_rows: int


class ImportError_(ValueError):
    """Raised when a file cannot be read at all."""


def load_records(filename: str, data: bytes) -> ImportResult:
    """Read an uploaded file's bytes into records. Raises ImportError_ on bad files."""
    lowered = filename.casefold()
    if lowered.endswith(".csv") or lowered.endswith(".txt"):
        header, rows = _read_csv(data)
    elif lowered.endswith(".xlsx") or lowered.endswith(".xlsm"):
        header, rows = _read_xlsx(data)
    elif lowered.endswith(".xls"):
        raise ImportError_(
            "Old .xls files are not supported. Open it in Excel and use Save As → .xlsx or .csv."
        )
    else:
        raise ImportError_("Upload a .csv or .xlsx file.")

    if not header:
        raise ImportError_("That file has no header row, so the columns cannot be identified.")

    mapping, warnings = _map_headers(header)
    if "customer_name" not in mapping:
        raise ImportError_(
            "No customer column was found. Add a column named 'Customer Name' "
            f"(columns seen: {', '.join(h for h in header if h) or 'none'})."
        )

    records: list[WorkRecord] = []
    skipped = 0
    for offset, raw in enumerate(rows):
        row_number = offset + 2  # +1 for the header, +1 because humans count from 1
        values = {field: _cell(raw, header, column) for field, column in mapping.items()}
        if not any(str(value).strip() for value in values.values()):
            skipped += 1
            continue
        records.append(_build_record(filename, row_number, values))
        if len(records) >= MAX_ROWS:
            warnings.append(f"Only the first {MAX_ROWS} rows were loaded.")
            break

    if not records:
        raise ImportError_("That file has a header row but no data rows.")

    return ImportResult(records, warnings, mapping, skipped)


def _build_record(source_label: str, row_number: int, values: Mapping[str, object]) -> WorkRecord:
    customer_name = _text(values.get("customer_name"))
    invoice_number = _text(values.get("invoice_number"))
    quantity = _money(values.get("quantity"))
    rate = _money(values.get("rate"))
    amount = _money(values.get("amount"))
    if quantity <= 0:
        quantity = Decimal("1")
    if not amount and rate:
        amount = (quantity * rate).quantize(Decimal("0.01"))
    return WorkRecord(
        source_id=deterministic_source_id(source_label, row_number, customer_name, invoice_number),
        row_number=row_number,
        customer_name=customer_name,
        job_name=_text(values.get("job_name")),
        company_name=_text(values.get("company_name")) or customer_name,
        email=_text(values.get("email")),
        phone=_text(values.get("phone")),
        invoice_number=invoice_number,
        invoice_date=_date(values.get("invoice_date")),
        item=_text(values.get("item")),
        description=_text(values.get("description")),
        quantity=quantity,
        rate=rate,
        amount=amount,
    )


# --- readers -----------------------------------------------------------------


def _read_csv(data: bytes) -> tuple[list[str], list[Sequence[object]]]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    rows = [row for row in reader]
    header_index = next((i for i, row in enumerate(rows) if any(str(c).strip() for c in row)), None)
    if header_index is None:
        return [], []
    header = [str(cell).strip() for cell in rows[header_index]]
    return header, rows[header_index + 1 :]


def _read_xlsx(data: bytes) -> tuple[list[str], list[Sequence[object]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the machine
        raise ImportError_(
            "Reading .xlsx needs the openpyxl library. Run: pip install openpyxl"
        ) from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportError_(f"That .xlsx file could not be opened: {exc}") from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    header_index = next(
        (i for i, row in enumerate(rows) if any(str(c).strip() for c in row if c is not None)),
        None,
    )
    if header_index is None:
        return [], []
    header = ["" if cell is None else str(cell).strip() for cell in rows[header_index]]
    return header, rows[header_index + 1 :]


# --- header matching ---------------------------------------------------------


def _normalise(header: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", header.casefold())).strip()


def _map_headers(header: Sequence[str]) -> tuple[dict[str, str], list[str]]:
    """Return {field: original column name} plus warnings about columns we ignored."""
    normalised = {_normalise(column): column for column in header if column}
    mapping: dict[str, str] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in normalised:
                mapping[field] = normalised[alias]
                break
    used = set(mapping.values())
    ignored = [column for column in header if column and column not in used]
    warnings: list[str] = []
    if ignored:
        warnings.append("Columns ignored: " + ", ".join(ignored) + ".")
    return mapping, warnings


def _cell(row: Sequence[object], header: Sequence[str], column: str) -> object:
    try:
        index = list(header).index(column)
    except ValueError:
        return ""
    return row[index] if index < len(row) else ""


# --- value cleaning ----------------------------------------------------------


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value).strip()


def _money(value: object) -> Decimal:
    """'$1,200.00' -> 1200.00, '(200)' -> -200, '' -> 0, '12%' -> 12."""
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = _text(value)
    if not text:
        return Decimal("0")
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.\-]", "", text)
    if cleaned in ("", "-", ".", "-."):
        return Decimal("0")
    try:
        number = Decimal(cleaned)
    except InvalidOperation:
        return Decimal("0")
    return -abs(number) if negative else number


def _date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    if not text:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None
